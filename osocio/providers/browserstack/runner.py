"""
runner.py
=========
Punto de entrada de BrowserStack para la app: lo llaman Envío de Leads, el CLI
(run.py --run-browserstack) y el runner autónomo. Corre un Excel de leads de un mercado
en un dispositivo de la nube y devuelve el resumen.

- Android y Mac: el motor de los navegadores (GenericCountryBase / BaseFormFiller) con el
  driver de BrowserStack. Heredan todas las reglas de la app (mapping, IDs dinámicos,
  '-' para omitir, documentos de Brasil, variantes CDC/T3, carpetas t1/t3, reintentos,
  columnas de resultado) sin duplicar nada.
- iPhone: el mismo motor (iphone_runner.FormFillerIphone), con el llenado por la capa
  nativa porque Safari en iPhone no deja entrar al iframe del form.

Mismo contrato de resumen que tenía lt_controller.run, así la UI no cambia de forma:
  pais, ok, failed, total, results_excel, video_url, session_id, error, fail_rows
"""

import os
import threading
from typing import Callable, Dict, Optional

from osocio.paths import DATA_DIR

PLATAFORMAS = ("iphone", "android", "mac")

# Sufijo del Excel de entrada por dispositivo (los de Mac/Android son los que ya
# generaba la app para LambdaTest, así siguen sirviendo).
SUFIJO_EXCEL = {"iphone": "Iphone", "android": "Android", "mac": "Mac"}

COLUMNA_VIDEO = "Video"


def excel_por_defecto(pais: str, plataforma: str, variante: str = "") -> str:
    return os.path.join(
        DATA_DIR, f"Lead_information_Formulario_{pais}_{SUFIJO_EXCEL[plataforma]}{variante or ''}.xlsx")


def nombre_build(pais: str, plataforma: str, programada: bool) -> str:
    # Sin guiones largos ni paréntesis: BrowserStack los descarta del nombre.
    disp = {"iphone": "iPhone", "android": "Android", "mac": "Mac"}[plataforma]
    modo = "Automatizado" if programada else "Envío Manual"
    return f"Osocio BS {disp} - {modo} - {pais}"


def _resumen_vacio(pais: str, error: str = "") -> Dict:
    return {"pais": pais, "ok": 0, "failed": 0, "total": 0, "results_excel": None,
            "video_url": "", "session_id": "", "error": error or None, "fail_rows": []}


def escribir_video(results_path: str, video_url: str, log: Callable = print) -> None:
    """Agrega/llena la columna Video en todas las filas y reordena el Excel."""
    if not (results_path and os.path.exists(results_path)):
        return
    try:
        from openpyxl import load_workbook
        wb = load_workbook(results_path)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        if COLUMNA_VIDEO in headers:
            col = headers.index(COLUMNA_VIDEO) + 1
        else:
            col = len(headers) + 1
            ws.cell(row=1, column=col).value = COLUMNA_VIDEO
        for fila in range(2, ws.max_row + 1):
            ws.cell(row=fila, column=col).value = video_url or "-"
        wb.save(results_path)
        from osocio.utils.excel_layout import reordenar_archivo
        reordenar_archivo(results_path, log=log)
    except Exception as e:
        log(f"  ⚠ No se pudo escribir la columna Video: {e}")


def run(pais: str, plataforma: str, excel_path: Optional[str] = None, variante: str = "",
        programada: bool = False, build_name: str = "", log_fn: Callable = print,
        stop_event: Optional[threading.Event] = None) -> Dict:
    """Corre el Excel de un mercado en un dispositivo de BrowserStack."""
    if plataforma not in PLATAFORMAS:
        return _resumen_vacio(pais, f"Plataforma de BrowserStack desconocida: {plataforma}")
    excel_path = excel_path or excel_por_defecto(pais, plataforma, variante)
    if not os.path.exists(excel_path):
        return _resumen_vacio(pais, f"Excel no encontrado: {excel_path}")
    build_name = build_name or nombre_build(pais, plataforma, programada)
    try:
        from osocio.providers.browserstack.bs_driver import load_credentials
        load_credentials()
    except Exception as e:
        return _resumen_vacio(pais, str(e))

    return _run_motor(pais, plataforma, excel_path, variante, programada, build_name, log_fn)


def _clase_motor(plataforma):
    """El motor de la app para cada dispositivo: el de los navegadores tal cual, salvo el
    iPhone, que es el mismo motor con el llenado por la capa nativa."""
    if plataforma == "iphone":
        from osocio.providers.browserstack.iphone_runner import FormFillerIphone
        return FormFillerIphone
    from osocio.core.generic_country_base import GenericCountryBase
    return GenericCountryBase


def _run_motor(pais, plataforma, excel_path, variante, programada, build_name, log_fn):
    from osocio.providers.browserstack.fabrica import FabricaNavegadorBrowserStack

    resumen = _resumen_vacio(pais)
    fab = FabricaNavegadorBrowserStack(plataforma, build_name, f"{pais} - {os.path.basename(excel_path)}",
                                       log=log_fn)
    form = _clase_motor(plataforma)(pais, browser=f"browserstack_{plataforma}", browser_factory=fab,
                              is_scheduled=programada, excel_suffix=variante, capturas=False)
    form.EXCEL_PATH = excel_path
    try:
        form.run()
    except Exception as e:
        resumen["error"] = str(e)
        log_fn(f"✗ Error en la corrida BrowserStack {plataforma}: {e}")
    finally:
        summ = getattr(form, "run_summary", None) or {}
        resumen.update(ok=int(summ.get("ok", 0)), failed=int(summ.get("fail", 0)),
                       total=int(summ.get("total", 0)), fail_rows=list(summ.get("fail_rows") or []),
                       results_excel=getattr(form, "RESULTADOS_PATH", None))
        if fab.sesion is not None:
            todo_ok = resumen["failed"] == 0 and resumen["ok"] > 0 and not resumen["error"]
            info = fab.sesion.cerrar(todo_ok, f"ok={resumen['ok']} fallidos={resumen['failed']}")
            resumen["session_id"] = fab.sesion.session_id
            resumen["video_url"] = info.get("video_url", "")
            if resumen["video_url"]:
                log_fn(f"  Video: {resumen['video_url']}")
        escribir_video(resumen["results_excel"], resumen["video_url"], log=log_fn)
    return resumen
