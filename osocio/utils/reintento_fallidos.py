"""
reintento_fallidos.py
=====================
Qué filas fallaron, leído de los Excels de resultados (no de memoria).

"Reintentar fallidos" usaba sólo lo que la app recordaba de la última corrida de ESTA
sesión: si la app se cerraba (o se reabría para tomar una actualización) el botón decía
que no había nada para reintentar. Ahora, para un país y dispositivo:

  1. Base: la última corrida COMPLETA (un resultado sin columna "Fila original"). Sus
     filas en rojo son las que fallaron, por número de fila.
  2. Los reintentos posteriores ajustan esa lista: su columna "Fila original" (la agrega
     el Excel temporal del reintento) dice a qué fila corresponde cada resultado. Si salió
     bien, sale de la lista; si sigue en rojo, queda.
  3. Se reintenta con lo que dice el Excel de entrada AHORA: si después de la corrida se
     corrigió una URL (o cualquier dato) de una fila que falló, va lo corregido.

Reintentos viejos sin "Fila original" (anteriores a este cambio) se ubican por URL; si
la URL se cambió después, simplemente no ajustan nada.
"""

import glob
import os
import re
from typing import List, Optional, Tuple

from osocio.core.variantes import basenames_de_corrida

ROJO_FALLA = "FFC7CE"   # el color con el que el motor pinta "Resultado" si el lead falló
COLUMNA_FILA_ORIGINAL = "Fila original"


def resultados_de(carpeta_resultados: str, pais: str, navegador: str, sufijo: str = "") -> List[str]:
    """Los Excels de resultados de esa corrida (resultados_<Pais>_<Disp>N.xlsx), del más
    viejo al más nuevo."""
    base = basenames_de_corrida(pais, navegador, False, sufijo)[0]
    patron = re.compile(re.escape(base) + r"\d+\.xlsx$", re.IGNORECASE)
    candidatos = [p for p in glob.glob(os.path.join(carpeta_resultados, "**", f"{base}*.xlsx"),
                                       recursive=True)
                  if patron.match(os.path.basename(p))]
    return sorted(candidatos, key=os.path.getmtime)


def ultimo_resultado(carpeta_resultados: str, pais: str, navegador: str,
                     sufijo: str = "") -> Optional[str]:
    """El Excel de resultados más reciente de esa corrida."""
    todos = resultados_de(carpeta_resultados, pais, navegador, sufijo)
    return todos[-1] if todos else None


def _norm(url) -> str:
    return str(url or "").strip().rstrip("/").lower()


def _leer(resultados_path: str):
    """Filas del resultado: [(fila, fila_original|None, en_rojo, URL, Formulario)] y si es
    de un reintento (tiene la columna "Fila original")."""
    from openpyxl import load_workbook
    ws = load_workbook(resultados_path).active
    hdr = [str(c.value or "").strip() for c in ws[1]]
    if "Resultado" not in hdr or "URL" not in hdr:
        return [], False
    col = {n: hdr.index(n) + 1 for n in ("Resultado", "URL", "Formulario", COLUMNA_FILA_ORIGINAL) if n in hdr}
    filas = []
    for fila in range(2, ws.max_row + 1):
        celda = ws.cell(row=fila, column=col["Resultado"])
        if not str(celda.value or "").strip():
            continue
        color = str(getattr(celda.fill.fgColor, "rgb", "") or "").upper()
        orig = None
        if COLUMNA_FILA_ORIGINAL in col:
            v = str(ws.cell(row=fila, column=col[COLUMNA_FILA_ORIGINAL]).value or "").strip()
            orig = int(float(v)) if v.replace(".", "", 1).isdigit() else None
        filas.append((fila, orig, color.endswith(ROJO_FALLA),
                      str(ws.cell(row=fila, column=col["URL"]).value or ""),
                      str(ws.cell(row=fila, column=col["Formulario"]).value or "") if "Formulario" in col else ""))
    return filas, COLUMNA_FILA_ORIGINAL in col


def urls_fallidas(resultados_path: str) -> List[Tuple[str, str]]:
    """(URL, Formulario) de cada fila con "Resultado" en rojo."""
    return [(u, f) for _, _, rojo, u, f in _leer(resultados_path)[0] if rojo]


def _filas_entrada_por_url(excel_entrada: str):
    """{(url, form): [filas]} del Excel de entrada, para ubicar reintentos viejos."""
    import pandas as pd
    df = pd.read_excel(excel_entrada, dtype=str).fillna("")
    cols = list(df.columns)
    c_url = "URL" if "URL" in cols else cols[0]
    c_form = "Formulario" if "Formulario" in cols else (cols[1] if len(cols) > 1 else None)
    idx = {}
    for i, r in df.iterrows():
        idx.setdefault((_norm(r[c_url]), _norm(r[c_form]) if c_form else ""), []).append(i + 2)
    return idx, len(df)


def filas_en_entrada(excel_entrada: str, fallidas: List[Tuple[str, str]]) -> List[int]:
    """Filas del Excel de entrada con esas URL + Formulario (cada falla, una fila distinta)."""
    if not fallidas or not os.path.exists(excel_entrada):
        return []
    idx, _ = _filas_entrada_por_url(excel_entrada)
    libres = {k: list(v) for k, v in idx.items()}
    out = []
    for url, form in fallidas:
        k = (_norm(url), _norm(form))
        cands = libres.get(k) or next((v for (u, _f), v in libres.items() if u == k[0] and v), None)
        if cands:
            out.append(cands.pop(0))
    return sorted(out)


def fallidas_desde_resultados(carpeta_resultados: str, pais: str, navegador: str,
                              excel_entrada: str, sufijo: str = "",
                              log=None) -> Tuple[Optional[str], List[int]]:
    """(Excel de resultados de la última corrida completa, filas del Excel de entrada a
    reintentar). Ver el docstring del módulo."""
    if not os.path.exists(excel_entrada):
        return None, []
    try:
        idx_url, n_entrada = _filas_entrada_por_url(excel_entrada)
        todos = [(p, *_leer(p)) for p in resultados_de(carpeta_resultados, pais, navegador, sufijo)]
        # Reintento = tiene "Fila original"; los de antes de ese cambio no la tienen, pero
        # tienen menos filas que el Excel de entrada.
        completos = [i for i, (_, filas, con_col) in enumerate(todos)
                     if not con_col and len(filas) >= n_entrada]
        if not completos:
            return None, []
        base = completos[-1]
        fallan = {fila for fila, _, rojo, _, _ in todos[base][1] if rojo}
        for _path, filas, _ in todos[base + 1:]:      # reintentos posteriores
            for _fila, orig, rojo, url, form in filas:
                if orig is None:                      # reintento viejo: por URL
                    cands = idx_url.get((_norm(url), _norm(form))) or []
                    if len(cands) != 1:
                        continue
                    orig = cands[0]
                (fallan.add if rojo else fallan.discard)(orig)
        filas = sorted(f for f in fallan if 2 <= f <= n_entrada + 1)
        if log is not None and len(filas) < len(fallan):
            log(f"[WARN] Reintentar fallidos: {len(fallan) - len(filas)} fila(s) que fallaron ya no "
                f"existen en {os.path.basename(excel_entrada)} (¿se borraron filas?).")
        return todos[base][0], filas
    except Exception as e:
        if log is not None:
            log(f"[WARN] Reintentar fallidos: no se pudieron leer los resultados de {pais}: {e}")
        return None, []
