"""Reintentar fallidos: filas en rojo de la última corrida completa del país/dispositivo,
ajustadas por los reintentos posteriores, con los datos del Excel de entrada de HOY."""
import os
import time

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from osocio.utils.reintento_fallidos import (
    COLUMNA_FILA_ORIGINAL, fallidas_desde_resultados, ultimo_resultado,
)

VERDE, ROJO, AMBAR = "C6EFCE", "FFC7CE", "FFC000"
L1, L2, L3 = "https://a.com/1", "https://a.com/2", "https://a.com/3"
F1, F2, F3 = "https://f.com/1", "https://f.com/2", "https://f.com/3"


def _resultado(carpeta, nombre, filas, fila_original=None):
    """filas: (URL, Formulario, color de 'Resultado'). fila_original: lista (reintento)."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Resultado", "Motivo", "URL", "Formulario"]
              + ([COLUMNA_FILA_ORIGINAL] if fila_original else []))
    for i, (url, form, color) in enumerate(filas):
        ws.append(["x", "-", url, form] + ([str(fila_original[i])] if fila_original else []))
        ws.cell(row=ws.max_row, column=1).fill = PatternFill(fill_type="solid", fgColor=color)
    path = os.path.join(carpeta, "t1", nombre)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    time.sleep(0.05)   # que el orden por fecha sea el de creación


def _entrada(path, filas):
    pd.DataFrame(filas, columns=["URL", "Formulario", "Nombre"]).to_excel(path, index=False)


def _correr(tmp_path, entrada, navegador="chrome"):
    return fallidas_desde_resultados(str(tmp_path / "resultados"), "Brasil", navegador, entrada)


def test_corrida_completa_toma_las_filas_en_rojo(tmp_path):
    entrada = str(tmp_path / "entrada.xlsx")
    _entrada(entrada, [(L1, F1, "Ana"), (L2, F2, "Bea"), (L3, F3, "Ceci")])
    # Ámbar = se envió por otro form: no se reintenta. Sólo el rojo.
    _resultado(str(tmp_path / "resultados"), "resultados_Brasil_BS_Mac1.xlsx",
               [(L1, F1, VERDE), (L2, F2, ROJO), (L3, F3, AMBAR)])
    usado, filas = _correr(tmp_path, entrada, "browserstack_mac")
    assert os.path.basename(usado) == "resultados_Brasil_BS_Mac1.xlsx"
    assert filas == [3]


def test_toma_los_cambios_del_excel_hechos_despues_de_la_corrida(tmp_path):
    entrada = str(tmp_path / "entrada.xlsx")
    _resultado(str(tmp_path / "resultados"), "resultados_Brasil_Chrome1.xlsx",
               [(L1, F1, VERDE), ("https://a.com/vieja", F2, ROJO)])
    _entrada(entrada, [(L1, F1, "Ana"), ("https://a.com/corregida", F2, "Bea")])
    _, filas = _correr(tmp_path, entrada)
    assert filas == [3]
    assert pd.read_excel(entrada).iloc[filas[0] - 2]["URL"] == "https://a.com/corregida"


def test_el_caso_de_brasil_url_cambiada_y_vuelta_a_poner_despues_de_un_reintento(tmp_path):
    """Corrida completa: fila 3 falla. Se cambia la URL y se reintenta (vuelve a fallar).
    Se pone la URL original de nuevo: el próximo reintento tiene que tomar la fila 3."""
    carpeta, entrada = str(tmp_path / "resultados"), str(tmp_path / "entrada.xlsx")
    _resultado(carpeta, "resultados_Brasil_BS_Mac1.xlsx",
               [(L1, F1, VERDE), ("https://sonic/acessorios/formulario", F2, ROJO), (L3, F3, VERDE)])
    _resultado(carpeta, "resultados_Brasil_BS_Mac2.xlsx",
               [("https://sonic/acessorios", F2, ROJO)], fila_original=[3])
    _entrada(entrada, [(L1, F1, "Ana"), ("https://sonic/acessorios/formulario", F2, "Bea"), (L3, F3, "Ceci")])
    usado, filas = _correr(tmp_path, entrada, "browserstack_mac")
    assert os.path.basename(usado) == "resultados_Brasil_BS_Mac1.xlsx"
    assert filas == [3]


def test_un_reintento_que_salio_bien_la_saca_de_la_lista(tmp_path):
    carpeta, entrada = str(tmp_path / "resultados"), str(tmp_path / "entrada.xlsx")
    _entrada(entrada, [(L1, F1, "Ana"), (L2, F2, "Bea"), (L3, F3, "Ceci")])
    _resultado(carpeta, "resultados_Brasil_Chrome1.xlsx", [(L1, F1, ROJO), (L2, F2, ROJO), (L3, F3, VERDE)])
    _resultado(carpeta, "resultados_Brasil_Chrome2.xlsx", [(L1, F1, VERDE), (L2, F2, ROJO)],
               fila_original=[2, 3])
    assert _correr(tmp_path, entrada)[1] == [3]
    # Una corrida completa nueva es la nueva base.
    _resultado(carpeta, "resultados_Brasil_Chrome3.xlsx", [(L1, F1, VERDE), (L2, F2, VERDE), (L3, F3, ROJO)])
    assert _correr(tmp_path, entrada)[1] == [4]


def test_reintento_viejo_sin_fila_original_se_ubica_por_url(tmp_path):
    carpeta, entrada = str(tmp_path / "resultados"), str(tmp_path / "entrada.xlsx")
    _entrada(entrada, [(L1, F1, "Ana"), (L2, F2, "Bea"), (L3, F3, "Ceci")])
    _resultado(carpeta, "resultados_Brasil_Chrome1.xlsx", [(L1, F1, ROJO), (L2, F2, ROJO), (L3, F3, VERDE)])
    _resultado(carpeta, "resultados_Brasil_Chrome2.xlsx", [(L1, F1, VERDE), (L2, F2, ROJO)])
    assert _correr(tmp_path, entrada)[1] == [3]


def test_no_mezcla_otro_dispositivo_ni_otra_variante(tmp_path):
    carpeta = str(tmp_path / "resultados")
    for nombre in ("resultados_Brasil_BS_Mac1.xlsx", "resultados_Brasil_T3_Chrome1.xlsx",
                   "resultados_Brasil_Chrome10.xlsx"):
        _resultado(carpeta, nombre, [])
    assert os.path.basename(ultimo_resultado(carpeta, "Brasil", "chrome")) == "resultados_Brasil_Chrome10.xlsx"
    assert ultimo_resultado(carpeta, "Brasil", "browserstack_iphone") is None


def test_el_excel_temporal_del_reintento_lleva_la_fila_original(tmp_path):
    """El Excel que arma la app para el reintento trae "Fila original" al final."""
    import importlib
    import inspect
    mi = importlib.import_module("osocio.interface.main_interface")
    assert "COLUMNA_FILA_ORIGINAL" in inspect.getsource(mi)
