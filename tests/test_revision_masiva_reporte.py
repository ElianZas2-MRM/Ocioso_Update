"""Reporte de la Revisión Masiva: columna Motivo, filas que no son URL y un Excel por
mercado a medida que terminan. Sin navegador: el revisor de cada fila es falso."""
import os

import openpyxl

from osocio.core import massive_check_runner as mcr
from osocio.core.massive_check_runner import form_insertado_texto, motivo_de_resultado

FORM = "https://secure-developments.com/commonwealth/colombia/gm_front/form/movs-visid-quote?model=n400"
OK_200 = "200 - ✅ PASS Valid link. HTTP Code: 200"


def _res(**kw):
    base = {"landing_status": OK_200, "form_coincide": "✅ PASS", "displayed": "✅ PASS",
            "secure_status": OK_200, "form_url_found": FORM, "error": ""}
    base.update(kw)
    return base


def test_motivo_ok():
    assert motivo_de_resultado(_res()) == "OK"


def test_motivo_landing():
    assert motivo_de_resultado(_res(
        landing_status="404 - ❌ FAIL Landing 404 / no encontrada (título: '404')",
        form_coincide="❌ FAIL No se pudo revisar el form: falló la landing",
        form_url_found="")) == "Landing 404 (no existe)"
    assert motivo_de_resultado(_res(landing_status="503 - ❌ FAIL HTTP Code: 503")).startswith(
        "Landing 503 (error del servidor)")
    m = motivo_de_resultado(_res(
        landing_status=("301 - 🔀 FAIL Link redirects. Code: 301, URL Original: https://a.com/x "
                        "redirects to: https://a.com/"),
        form_coincide=f"❌ FAIL Mismatch: expected {FORM} but found ningún form insertado",
        form_url_found=""))
    assert m == "Landing redirige (301) a https://a.com/"
    assert motivo_de_resultado(_res(landing_status="Timeout - ❌ FAIL La URL no respondió en 8s")
                               ).startswith("Landing no respondió (timeout)")
    # Un CTA que no lleva al form no es un 404.
    m = motivo_de_resultado(_res(landing_status="CTA - ❌ FAIL No se encontró ningún CTA (a.x) en https://a.com",
                                 form_coincide="❌ FAIL No se pudo revisar el form: falló la landing"))
    assert m.startswith("CTA de la landing: No se encontró ningún CTA") and "404" not in m


def test_motivo_form():
    otro = "https://secure-developments.com/commonwealth/colombia/gm_front/form/otro"
    assert motivo_de_resultado(_res(
        form_coincide=f"❌ FAIL Mismatch: expected {FORM} but found {otro}",
        form_url_found=otro, displayed="❌ FAIL")) == f"Form distinto al esperado: en la landing está {otro}"  # sin esperado
    assert motivo_de_resultado(_res(
        form_coincide=f"❌ FAIL Mismatch: expected {FORM} but found ningún form insertado",
        form_url_found="", displayed="❌ FAIL")) == "La landing no tiene el form insertado"
    assert motivo_de_resultado(_res(displayed="❌ FAIL")) == "Form correcto pero no se ven sus campos"
    assert motivo_de_resultado(_res(displayed="❌ FAIL", secure_status="404 - ❌ FAIL HTTP Code: 404")) == (
        "El form insertado es el correcto pero da 404 (no carga)")
    assert motivo_de_resultado(_res(secure_status="404 - ❌ FAIL HTTP Code: 404")) == (
        "La URL del form del Excel da 404")


def test_motivo_dice_cual_se_esperaba():
    otro = "https://secure-developments.com/commonwealth/colombia/gm_front/form/otro"
    assert motivo_de_resultado(_res(
        form_coincide=f"❌ FAIL Mismatch: expected {FORM} but found {otro}",
        form_url_found=otro, displayed="❌ FAIL"), FORM) == (
        "Form distinto al esperado (ruta distinta: está 'otro' y se esperaba 'movs-visid-quote'; "
        f"falta el parámetro model=n400): en la landing está {otro} (se esperaba {FORM})")
    assert motivo_de_resultado(_res(
        form_coincide=f"❌ FAIL Mismatch: expected {FORM} but found ningún form insertado",
        form_url_found="", displayed="❌ FAIL"), FORM) == (
        f"La landing no tiene el form insertado (se esperaba {FORM})")


def test_iframe_vacio_y_redireccion_en_el_navegador():
    vacio = _res(form_coincide=f"❌ FAIL Mismatch: expected {FORM} but found un iframe vacío, sin URL",
                 form_url_found="", displayed="❌ FAIL",
                 iframes_vacios=["financiamiento rural reemplazo form"])
    assert motivo_de_resultado(vacio, FORM) == (
        "La landing tiene un iframe vacío, sin URL ('financiamiento rural reemplazo form') "
        f"(se esperaba {FORM})")
    assert form_insertado_texto(vacio) == "— (un iframe vacío, sin URL ('financiamiento rural reemplazo form'))"

    js = _res(landing_status="JS - 🔀 FAIL La landing redirige en el navegador a https://www.chevrolet.com.br/",
              form_coincide=f"❌ FAIL Mismatch: expected {FORM} but found ningún form insertado",
              form_url_found="", displayed="❌ FAIL")
    assert motivo_de_resultado(js, FORM) == "La landing redirige en el navegador a https://www.chevrolet.com.br/"
    assert form_insertado_texto(js) == "— (falló la landing: ver Motivo)"
    assert form_insertado_texto(_res(form_url_found="")) == "— (ningún form insertado)"
    assert form_insertado_texto(_res()) == FORM


def test_misma_pagina():
    assert mcr._misma_pagina("https://chevrolet.com.br/x/", "https://www.chevrolet.com.br/x?a=1#b")
    assert not mcr._misma_pagina("https://www.chevrolet.com.br/servicos/blindagem", "https://www.chevrolet.com.br/")


def _matriz(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "COLOMBIA"
    ws.append(["SEGMENTO", "ESTADO", "URL LIVE", "URL SECURE", "Comentarios"])
    # Como en la matriz real: celdas con link (Colombia muestra el título y la URL va en el link).
    ws.append(["N400", "ON", "N400 Cargo | Chevrolet", FORM, ""])
    ws["C2"].hyperlink = "https://www.chevrolet.com.co/vans/n400-cargo"
    ws["D2"].hyperlink = FORM
    ws.append(["Tabla", "ON", "LIVE URL", "FORMS PROD", ""])           # otra tabla en la hoja
    ws.append(["Dia", "ON", "Lunes - ANDROID", "", ""])
    ws.append(["Viejo", "OFF", "https://www.chevrolet.com.co/x", FORM, ""])
    ws2 = wb.create_sheet("CHILE")
    ws2.append(["SEGMENTO", "ESTADO", "URL LIVE", "URL SECURE", "Comentarios"])
    ws2.append(["Solo form", "ON", "Lo usa GM por codigo, no existe en una lading",
                "https://secure-developments.com/commonwealth/chile/gm_forms/vehicle-shopper", ""])
    wb.save(path)


class _FillerFalso:
    revisadas = []

    def __init__(self, country_name, **kw):
        self.driver = self
        self.screenshot_manager = None

    def quit(self):
        pass

    def setup_directories_and_files(self, **kw):
        pass

    def initialize_browser(self):
        pass

    def check_single_url(self, landing, form, row, tomar_capturas=True):
        _FillerFalso.revisadas.append((landing, form))
        return {"ok": True, "current_url": landing or form, "landing_status": OK_200,
                "form_url_found": form, "form_coincide": "✅ PASS", "displayed": "✅ PASS",
                "secure_status": OK_200, "forms_count": 1, "error": ""}


def _tabla(ws):
    """Tabla de todas las filas de la pestaña del mercado: {fila original: {columna: valor}}."""
    filas = list(ws.iter_rows(values_only=True))
    i = next(n for n, f in enumerate(filas) if f[0] == "Fila" and "currentUrl" in f)
    cab = filas[i]
    return cab, {f[0]: dict(zip(cab, f)) for f in filas[i + 1:] if f[0] is not None}


def test_excel_por_mercado_y_filas_sin_url(tmp_path, monkeypatch):
    import osocio.paths
    monkeypatch.setattr(osocio.paths, "BASE_DIR", str(tmp_path))
    monkeypatch.setattr(mcr, "MassiveFormFiller", _FillerFalso)
    _FillerFalso.revisadas = []
    matriz = tmp_path / "Matriz.xlsx"
    _matriz(matriz)
    cols = {"segmento": "SEGMENTO", "estado": "ESTADO", "url_live": "URL LIVE", "url_secure": "URL SECURE"}
    avisos = []

    ok, info = mcr.run_massive_check(str(matriz), cols, ["Colombia", "Chile"], tomar_capturas=False,
                                     progress_callback=lambda p, m, pct: avisos.append((p, m)),
                                     paralelo=True)
    assert ok, info

    # Solo se revisan las filas con URL; la de Chile, como form suelto.
    assert sorted(_FillerFalso.revisadas) == sorted([
        ("https://www.chevrolet.com.co/vans/n400-cargo", FORM),
        ("", "https://secure-developments.com/commonwealth/chile/gm_forms/vehicle-shopper"),
    ])
    assert info["total_failed"] == 0 and info["total_skipped"] == 1

    # Un Excel por mercado, con solo su hoja, y el aviso de que está listo.
    por_mercado = sorted(os.path.basename(p) for p in info["reportes_por_mercado"])
    assert [n.rsplit("_", 1)[-1] for n in por_mercado] == ["CHILE.xlsx", "COLOMBIA.xlsx"]
    for p in info["reportes_por_mercado"]:
        assert os.path.dirname(p).endswith("por_mercado")
        assert len(openpyxl.load_workbook(p).sheetnames) == 1
    assert ("Colombia", "Listo — Excel del mercado guardado") in avisos

    # El completo: una pestaña por mercado y el resumen. Abajo de las fallas, todas las
    # filas con su número de fila del Excel original y en qué quedaron.
    wb_todo = openpyxl.load_workbook(info["reporte_path"])
    assert wb_todo.sheetnames == ["COLOMBIA", "CHILE", mcr.RESUMEN_SHEET_NAME]
    ws = wb_todo["COLOMBIA"]
    cab, tabla = _tabla(ws)
    assert list(cab[:3]) == ["Fila", "Motivo", "currentUrl"]
    assert tabla[2]["Motivo"] == "OK"
    assert tabla[5]["Motivo"] == "Omitido (OFF)"
    assert tabla[3]["Motivo"] == "No se revisa (sin URL)"     # "LIVE URL": otra tabla

    # Sin links clickeables; la URL que solo estaba en el link queda escrita.
    for p in [info["reporte_path"]] + info["reportes_por_mercado"]:
        for hoja in openpyxl.load_workbook(p).worksheets:
            assert not [c.coordinate for fila in hoja.iter_rows() for c in fila if c.hyperlink]
    assert tabla[2]["URL LIVE"] == "https://www.chevrolet.com.co/vans/n400-cargo"


def test_reporte_viejo_sin_motivo_no_duplica_el_bloque(tmp_path, monkeypatch):
    import osocio.paths
    monkeypatch.setattr(osocio.paths, "BASE_DIR", str(tmp_path))
    monkeypatch.setattr(mcr, "MassiveFormFiller", _FillerFalso)
    matriz = tmp_path / "Viejo.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "COLOMBIA"
    ws.append(mcr.RESULT_HEADERS[1:] + ["SEGMENTO", "ESTADO", "URL LIVE", "URL SECURE", "Comentarios"])
    ws.append([""] * 7 + ["N400", "ON", "https://www.chevrolet.com.co/vans/n400-cargo", FORM, ""])
    wb.save(matriz)
    cols = {"segmento": "SEGMENTO", "estado": "ESTADO", "url_live": "URL LIVE", "url_secure": "URL SECURE"}

    ok, info = mcr.run_massive_check(str(matriz), cols, ["Colombia"], tomar_capturas=False)
    assert ok, info
    cab, _ = _tabla(openpyxl.load_workbook(info["reporte_path"])["COLOMBIA"])
    assert list(cab[1:9]) == mcr.RESULT_HEADERS and cab.count("currentUrl") == 1


def test_form_aem_por_ruta_se_revisa(tmp_path, monkeypatch):
    """La ruta de un form AEM (/content/forms/af/...) es un form esperado válido; sola, sin
    landing, no hay nada que abrir."""
    import osocio.paths
    monkeypatch.setattr(osocio.paths, "BASE_DIR", str(tmp_path))
    monkeypatch.setattr(mcr, "MassiveFormFiller", _FillerFalso)
    _FillerFalso.revisadas = []
    ruta = "/content/forms/af/cadillac/br/solicitar-contato"
    matriz = tmp_path / "Aem.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BRASIL"
    ws.append(["SEGMENTO", "ESTADO", "URL LIVE", "URL SECURE", "Comentarios"])
    ws.append(["Cadillac", "ON", "https://www.cadillac.com.br/contato", ruta, ""])
    ws.append(["Solo ruta", "ON", "", ruta, ""])
    wb.save(matriz)
    cols = {"segmento": "SEGMENTO", "estado": "ESTADO", "url_live": "URL LIVE", "url_secure": "URL SECURE"}
    ok, info = mcr.run_massive_check(str(matriz), cols, ["Brasil"], tomar_capturas=False)
    assert ok, info
    assert _FillerFalso.revisadas == [("https://www.cadillac.com.br/contato", ruta)]


def test_celdas_combinadas_no_tapan_el_motivo(tmp_path, monkeypatch):
    """BRASIL tiene una tabla REGRESSIONS con celdas combinadas en la columna A: al insertar
    el bloque de resultados se tienen que correr con el resto, no tapar la columna Motivo."""
    import osocio.paths
    monkeypatch.setattr(osocio.paths, "BASE_DIR", str(tmp_path))
    monkeypatch.setattr(mcr, "MassiveFormFiller", _FillerFalso)
    matriz = tmp_path / "Combinadas.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BRASIL"
    ws.append(["SEGMENTO", "ESTADO", "URL LIVE", "URL SECURE", "Comentarios"])
    ws.append(["REGRESSIONS", "", "", "", ""])
    ws.merge_cells("A2:D2")
    for _ in range(3):
        ws.append(["", "LUNES", "https://www.chevrolet.com.br/solicitar-contato", FORM, ""])
    ws.merge_cells("A3:A5")
    wb.save(matriz)
    cols = {"segmento": "SEGMENTO", "estado": "ESTADO", "url_live": "URL LIVE", "url_secure": "URL SECURE"}

    ok, info = mcr.run_massive_check(str(matriz), cols, ["Brasil"], tomar_capturas=False)
    assert ok, info
    _, tabla = _tabla(openpyxl.load_workbook(info["reporte_path"])["BRASIL"])
    assert [tabla[r]["Motivo"] for r in (3, 4, 5)] == ["OK"] * 3
    assert tabla[2]["SEGMENTO"] == "REGRESSIONS"


def test_iframe_que_no_es_form():
    assert mcr._es_url_de_form("https://secure-developments.com/commonwealth/brasil/gm_forms/raq")
    assert mcr._es_url_de_form("https://secure-developments.com/commonwealth/chile/gm_frontend/x")
    assert mcr._es_url_de_form("https://www.chevrolet.com.br/request-proposal/jcr:content/root/form/formcontainer.iframe.br")
    assert not mcr._es_url_de_form("https://secure-developments.com/commonwealth/brasil/visualizer_sonic/")


def test_mismo_form_exacto_y_que_cambia():
    esp = "https://secure-developments.com/commonwealth/brasil/gm_forms/tenho-interese-trailblazer?model=trailblazer"
    assert mcr._mismo_form(esp + " ", esp)
    casos = {
        "https://secure-developments.com/commonwealth/brasil/gm_forms/myr2026?model=trailblazer":
            (True, ["ruta distinta: está 'myr2026' y se esperaba 'tenho-interese-trailblazer'"]),
        esp.replace("https://", "https://www."):
            (False, ["el insertado tiene 'www.' y el esperado no"]),
        esp.replace("https://", "http://"):
            (False, ["protocolo: está http:// y se esperaba https://"]),
        esp.replace("?model", "/?model"):
            (False, ["barra final: el insertado termina en '/' y el esperado no"]),
        esp.replace("=trailblazer", "=onix"):
            (False, ["model distinto: está 'onix' y se esperaba 'trailblazer'"]),
        esp.replace("=trailblazer", "=Trailblazer"):
            (False, ["model escrito distinto: está 'Trailblazer' y se esperaba 'trailblazer'"]),
        esp + "&utm=x": (False, ["parámetro de más: utm=x"]),
        esp.replace("=trailblazer", "=trailblazer%C2%A0"):
            (False, ["model tiene un espacio invisible (espacio no separable) de más: está 'trailblazer ' "
                     "y se esperaba 'trailblazer'"]),
        esp + "#form": (False, ["al insertado le sobra '#form'"]),
        esp.split("?")[0]: (False, ["falta el parámetro model=trailblazer"]),
        "https://otro.com/commonwealth/brasil/gm_forms/tenho-interese-trailblazer?model=trailblazer":
            (True, ["dominio distinto: está otro.com y se esperaba secure-developments.com"]),
    }
    for insertado, esperado_dif in casos.items():
        assert not mcr._mismo_form(insertado, esp)
        assert mcr.diferencias_form(insertado, esp) == esperado_dif, insertado


def test_motivo_form_parecido_dice_que_cambia():
    esp = "https://secure-developments.com/commonwealth/argentina/gm_front/form/movs-pda-visid?model=tracker"
    ins = esp.replace("=tracker", "=onix")
    m = motivo_de_resultado(_res(form_coincide=f"❌ FAIL Mismatch: expected {esp} but found {ins}",
                                 form_url_found=ins), esp)
    assert m == (f"Form casi igual al esperado pero no idéntico (model distinto: está 'onix' y se esperaba "
                 f"'tracker'): en la landing está {ins} (se esperaba {esp})")


def test_tipo_de_falla():
    assert mcr.tipo_de_falla("Landing 404 (no existe)")[0] == "404"
    assert mcr.tipo_de_falla("Landing redirige (301) a https://a.com/")[0] == "redirige"
    assert mcr.tipo_de_falla("La landing redirige en el navegador a https://a.com/")[0] == "redirige"
    assert mcr.tipo_de_falla("Landing 503 (error del servidor)")[0] == "no_carga"
    assert mcr.tipo_de_falla("La landing tiene un iframe vacío, sin URL")[0] == "sin_form"
    assert mcr.tipo_de_falla("El form insertado es el correcto pero da 404 (no carga)")[0] == "form_caido"
    assert mcr.tipo_de_falla("Form distinto al esperado (ruta distinta: ...)")[2] == "no_coincide"
    assert mcr.tipo_de_falla("Form casi igual al esperado pero no idéntico (...)")[0] == "casi_igual"
    assert mcr.tipo_de_falla("Error navegación: x")[0] == "otros"


class _FillerConFallas(_FillerFalso):
    def check_single_url(self, landing, form, row, tomar_capturas=True):
        if "rota" in landing:
            return {"ok": False, "current_url": landing, "landing_status": "404 - ❌ FAIL HTTP Code: 404",
                    "form_url_found": "", "form_coincide": "❌ FAIL No se pudo revisar el form: falló la landing",
                    "displayed": "❌ FAIL", "secure_status": OK_200, "forms_count": 0, "error": "404"}
        if "otra" in landing:
            otro = FORM.replace("movs-visid-quote", "otro")
            return {"ok": False, "current_url": landing, "landing_status": OK_200, "form_url_found": otro,
                    "form_coincide": f"❌ FAIL Mismatch: expected {FORM} but found {otro}",
                    "displayed": "✅ PASS", "secure_status": OK_200, "forms_count": 1, "error": ""}
        return super().check_single_url(landing, form, row, tomar_capturas)


def test_pestana_de_mercado_fallas_por_tipo_y_todas_las_filas(tmp_path, monkeypatch):
    import osocio.paths
    monkeypatch.setattr(osocio.paths, "BASE_DIR", str(tmp_path))
    monkeypatch.setattr(mcr, "MassiveFormFiller", _FillerConFallas)
    matriz = tmp_path / "Fallas.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "COLOMBIA"
    ws.append(["SEGMENTO", "ESTADO", "URL LIVE", "URL SECURE", "Comentarios"])
    ws.append(["Bien", "ON", "https://www.chevrolet.com.co/bien", FORM, ""])
    ws.append(["Rota", "ON", "https://www.chevrolet.com.co/rota", FORM, ""])
    ws.append(["Otra", "ON", "https://www.chevrolet.com.co/otra", FORM, ""])
    ws.append(["Vieja", "OFF", "https://www.chevrolet.com.co/vieja", FORM, ""])
    ws2 = wb.create_sheet("CHILE")
    ws2.append(["SEGMENTO", "ESTADO", "URL LIVE", "URL SECURE", "Comentarios"])
    ws2.append(["Bien", "ON", "https://www.chevrolet.cl/bien", FORM, ""])
    wb.save(matriz)
    cols = {"segmento": "SEGMENTO", "estado": "ESTADO", "url_live": "URL LIVE", "url_secure": "URL SECURE"}
    ok, info = mcr.run_massive_check(str(matriz), cols, ["Colombia", "Chile"], tomar_capturas=False)
    assert ok, info

    def textos(ws):
        return [str(c.value) for fila in ws.iter_rows() for c in fila if c.value is not None]

    por_mercado = {os.path.basename(p).rsplit("_", 1)[-1]: p for p in info["reportes_por_mercado"]}
    for ws in (openpyxl.load_workbook(info["reporte_path"])["COLOMBIA"],
               openpyxl.load_workbook(por_mercado["COLOMBIA.xlsx"])["COLOMBIA"]):
        t = textos(ws)
        assert "Revisadas: 3 · OK: 1 · FAIL: 2 · Omitidas (OFF): 1 · No se revisan (sin URL): 0" in t
        # Arriba las fallas por tipo: primero lo que hay que reportar, después lo que no coincide.
        assert t.index("Landing da 404 (1)") < t.index(
            "El form de la landing es otro (no coincide con la matriz) (1)") < t.index(
            "Todas las filas (Fila = número de fila en el Excel original)")
        # Abajo todas las filas, con su fila original.
        _, tabla = _tabla(ws)
        assert {f: tabla[f]["Motivo"].split(" ")[0] for f in tabla} == {
            2: "OK", 3: "Landing", 4: "Form", 5: "Omitido"}
    assert "Sin fallas ✔" in textos(openpyxl.load_workbook(por_mercado["CHILE.xlsx"])["CHILE"])


def test_el_reporte_nuevo_lo_leen_solo_fails_y_el_mail(tmp_path, monkeypatch):
    import osocio.paths
    from osocio.interface.helpers_interface import analizar_resultados_masivos
    monkeypatch.setattr(osocio.paths, "BASE_DIR", str(tmp_path))
    monkeypatch.setattr(mcr, "MassiveFormFiller", _FillerConFallas)
    matriz = tmp_path / "Lectores.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "COLOMBIA"
    ws.append(["SEGMENTO", "ESTADO", "URL LIVE", "URL SECURE", "Comentarios"])
    ws.append(["Bien", "ON", "https://www.chevrolet.com.co/bien", FORM, ""])
    ws.append(["Rota", "ON", "https://www.chevrolet.com.co/rota", FORM, ""])
    wb.save(matriz)
    cols = {"segmento": "SEGMENTO", "estado": "ESTADO", "url_live": "URL LIVE", "url_secure": "URL SECURE"}
    ok, info = mcr.run_massive_check(str(matriz), cols, ["Colombia"], tomar_capturas=False)
    assert ok, info
    previo = openpyxl.load_workbook(info["reporte_path"], data_only=True)
    assert set(mcr.load_passed_results_from_wb(previo, cols)) == {
        ("https://www.chevrolet.com.co/bien", FORM.lower())}
    assert analizar_resultados_masivos(info["reporte_path"])["market_summaries"]["COLOMBIA"] == {
        "total": 2, "pass": 1, "fail": 1}


class _Resp:
    status_code, history, url = 200, [], "https://www.chevrolet.com.ec/x"


def _filler_http(monkeypatch, fallos):
    """MassiveFormFiller sin navegador: solo para _http_status_str."""
    import requests
    llamadas = {"n": 0}

    def get(url, **kw):
        llamadas["n"] += 1
        if llamadas["n"] <= fallos:
            raise requests.exceptions.ConnectionError("('Connection aborted.', RemoteDisconnected('x'))")
        return _Resp()
    monkeypatch.setattr(mcr.requests, "get", get)
    monkeypatch.setattr(mcr.time, "sleep", lambda s: None)
    return object.__new__(mcr.MassiveFormFiller), llamadas


def test_corte_de_conexion_se_reintenta(monkeypatch):
    f, llamadas = _filler_http(monkeypatch, fallos=2)
    assert f._http_status_str("https://www.chevrolet.com.ec/x")[0].startswith("200 - ✅ PASS")
    assert llamadas["n"] == 3


def test_corte_de_conexion_que_sigue_es_fail_y_lo_dice(monkeypatch):
    f, llamadas = _filler_http(monkeypatch, fallos=99)
    status = f._http_status_str("https://www.chevrolet.com.ec/x")[0]
    assert llamadas["n"] == 3 and "ConnectionError (tras 3 intentos)" in status
    assert motivo_de_resultado(_res(landing_status=status)).startswith(
        "Landing inaccesible: ConnectionError (tras 3 intentos)")
