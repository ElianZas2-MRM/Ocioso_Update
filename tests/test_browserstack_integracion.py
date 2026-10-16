"""BrowserStack integrado a la app, sin red.

El riesgo de esta integración no es que algo explote al importar: es que un dispositivo
quede mal cableado entre la UI, el despacho y los resultados (y el mail salga sin
adjunto), o que el iPhone —que llena por la capa nativa— se aparte de las reglas de datos
de la app. Estos tests vigilan eso.
"""
import json
import os

import pytest

from osocio.core.campos_inferencia import ids_para_etiqueta
from osocio.core.variantes import basenames_de_corrida
from osocio.providers.browserstack import dispositivos, runner
from osocio.providers.browserstack.ios_nativo import OMITIR


# --- tabla única de dispositivos ----------------------------------------------------

def test_los_tres_dispositivos_y_sus_caras():
    assert [d.etiqueta for d in dispositivos.DISPOSITIVOS_BS] == ["iPhone BS", "Android BS", "Mac BS"]
    iphone = dispositivos.por_cualquier_nombre("iPhone BS")
    for nombre in ("iphone bs", "iphone", "bs_iphone", "browserstack_iphone"):
        assert dispositivos.por_cualquier_nombre(nombre) is iphone
    assert iphone.sufijo_excel == "Iphone"
    assert iphone.token_email == "bsiphone"


def test_un_navegador_local_no_es_de_la_nube():
    for nombre in ("chrome", "firefox", "edge", "", None):
        assert not dispositivos.es_browserstack(nombre)


def test_la_seleccion_guardada_con_lambdatest_pasa_a_browserstack():
    viejo = {"chrome": True, "mac lt": True, "android lt": False}
    assert dispositivos.migrar_seleccion(viejo) == {"chrome": True, "mac bs": True,
                                                    "android bs": False}


@pytest.mark.parametrize("configurado,esperado", [(None, 3), (1, 1), (5, 5), (50, 10), (0, 1)])
def test_tope_de_sesiones_en_paralelo(monkeypatch, tmp_path, configurado, esperado):
    cfg = {} if configurado is None else {"browserstack": {"max_paralelo": configurado}}
    (tmp_path / "config_global.json").write_text(json.dumps(cfg), encoding="utf-8")
    monkeypatch.setattr("osocio.paths.JSON_DIR", str(tmp_path))
    assert dispositivos.max_paralelo() == esperado


def test_el_excel_por_dispositivo_es_el_que_usa_la_ui():
    from osocio.interface.main_interface import _device_excel_suffix
    assert _device_excel_suffix("iPhone BS") == "Iphone"
    assert _device_excel_suffix("Android BS") == "Android"
    assert _device_excel_suffix("Mac BS") == "Mac"
    assert _device_excel_suffix("Chrome") == "Chrome"
    assert runner.excel_por_defecto("Brasil", "mac", "_T3").endswith(
        "Lead_information_Formulario_Brasil_Mac_T3.xlsx")


# --- resultados: mismas carpetas y nombres que los navegadores ------------------------

@pytest.mark.parametrize("navegador,sufijo,programada,esperado", [
    ("browserstack_mac", "", False, "resultados_Brasil_BS_Mac"),
    ("browserstack_iphone", "_CDC", False, "resultados_Brasil_CDC_BS_iPhone"),
    ("browserstack_android", "_T3", True, "Automatizacion_Brasil_T3_BS_Android"),
])
def test_nombre_de_los_resultados(navegador, sufijo, programada, esperado):
    excel, _capturas = basenames_de_corrida("Brasil", navegador, programada=programada, sufijo=sufijo)
    assert excel == esperado


def _solo_caracteres_que_acepta_browserstack(nombre):
    # Letras sin acento, dígitos, espacio y . : - [ ] / @ & ' _ (lo demás lo descarta).
    import re
    return re.fullmatch(r"[A-Za-z0-9 .:\-\[\]/@&'_]+", nombre) is not None


def test_un_build_por_dia():
    from datetime import datetime
    nombre = runner.nombre_build_del_dia(False, datetime(2026, 10, 14, 11, 9))
    assert nombre == "Osocio BS - Envio Manual - Miercoles 2026-10-14"
    assert _solo_caracteres_que_acepta_browserstack(nombre)
    # Otra corrida el mismo día (todos los mercados) -> mismo build.
    assert runner.nombre_build_del_dia(False, datetime(2026, 10, 14, 17, 45)) == nombre
    # Otro día -> otro build; las programadas van aparte.
    assert runner.nombre_build_del_dia(False, datetime(2026, 10, 15, 9, 0)) == \
        "Osocio BS - Envio Manual - Jueves 2026-10-15"
    assert runner.nombre_build_del_dia(True, datetime(2026, 10, 14)).startswith("Osocio BS - Automatizado")


def test_dentro_del_build_cada_sesion_dice_su_dispositivo():
    sesion = runner.nombre_sesion("Colombia", "iphone", "data/Lead_information_Formulario_Colombia_Iphone.xlsx")
    assert sesion == "iPhone - Colombia - Lead_information_Formulario_Colombia_Iphone.xlsx"
    assert _solo_caracteres_que_acepta_browserstack(runner.nombre_build("Colombia", "mac", False))


# --- sin capturas en la nube ----------------------------------------------------------

class _DriverFalso:
    name = "safari"


def test_sin_capturas_pero_con_la_senal_de_evidencia(tmp_path):
    from osocio.core.screenshot_manager import ScreenshotManagerNulo
    gestor = ScreenshotManagerNulo(_DriverFalso(), str(tmp_path))
    assert gestor.fname("form", "typage", 1) == ScreenshotManagerNulo.TEXTO_EVIDENCIA
    assert gestor.take_form_screenshot(1, "typage", full_page=True) is None
    assert gestor.take_landing_screenshot(1, "inicial") is None
    assert list(tmp_path.iterdir()) == []


# --- contrato del runner ----------------------------------------------------------------

def test_runner_plataforma_desconocida():
    res = runner.run("Brasil", "windows")
    assert res["error"] and res["ok"] == 0 and res["fail_rows"] == []


def test_runner_sin_excel(tmp_path):
    res = runner.run("Brasil", "mac", excel_path=str(tmp_path / "no_existe.xlsx"))
    assert "Excel no encontrado" in res["error"]


def test_runner_sin_credenciales_no_abre_nada(monkeypatch, tmp_path):
    excel = tmp_path / "leads.xlsx"
    excel.write_bytes(b"x")

    def _sin_credenciales():
        raise FileNotFoundError("No se encontraron credenciales de BrowserStack.")
    monkeypatch.setattr("osocio.providers.browserstack.bs_driver.load_credentials", _sin_credenciales)
    res = runner.run("Brasil", "iphone", excel_path=str(excel))
    assert "credenciales" in res["error"]
    assert set(res) >= {"pais", "ok", "failed", "total", "results_excel", "video_url",
                        "session_id", "error", "fail_rows"}


def test_cada_dispositivo_corre_el_motor_de_la_app():
    from osocio.core.generic_country_base import GenericCountryBase
    from osocio.providers.browserstack.iphone_runner import FormFillerIphone
    assert runner._clase_motor("mac") is GenericCountryBase
    assert runner._clase_motor("android") is GenericCountryBase
    assert runner._clase_motor("iphone") is FormFillerIphone
    assert issubclass(FormFillerIphone, GenericCountryBase)


def test_la_columna_video_se_escribe_en_todas_las_filas(tmp_path):
    from openpyxl import Workbook, load_workbook
    path = tmp_path / "resultados.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Resultado", "URL"])
    ws.append(["Lead enviado correctamente", "https://a"])
    ws.append(["Error", "https://b"])
    wb.save(path)
    runner.escribir_video(str(path), "https://video", log=lambda *_: None)
    ws = load_workbook(path).active
    col = [c.value for c in ws[1]].index("Video") + 1
    assert [ws.cell(row=r, column=col).value for r in (2, 3)] == ["https://video"] * 2


# --- iPhone: etiqueta visible -> id del mapping ------------------------------------------

@pytest.mark.parametrize("etiqueta,id_esperado", [
    ("Nombre*", "firstname"), ("Nome", "firstname"), ("Sobrenome", "lastname"),
    ("Cédula", "document"), ("CPF", "document"), ("CEP", "cep"),
    ("Tipo de documento", "document-type"), ("Concessionária", "dealer"),
    ("Departamento", "region"), ("Cidade", "city"), ("E-mail", "email"),
    ("Data estimada de compra", "estimated-day"), ("Modelos", "models"),
])
def test_etiqueta_a_id(etiqueta, id_esperado):
    assert id_esperado in ids_para_etiqueta(etiqueta)


@pytest.mark.parametrize("etiqueta", ["Acepto los términos y condiciones", "Username", ""])
def test_etiquetas_que_no_son_campos_del_mapping(etiqueta):
    # "Acepto" contiene "cep": las claves cortas van por palabra entera.
    assert ids_para_etiqueta(etiqueta) == ()


# --- iPhone: las reglas de datos son las de la app ---------------------------------------

@pytest.fixture
def iphone_brasil(monkeypatch):
    from osocio.providers.browserstack.iphone_runner import FormFillerIphone
    form = FormFillerIphone("Brasil", browser="browserstack_iphone", capturas=False)
    form.begin_row_tracking()
    form._tipo_doc_ios = ""
    monkeypatch.setattr(form, "_cargar_ids_dinamicos", lambda: {"lastname": ["Fijo"]})
    return form


def _form_data(form, valores_por_nombre):
    """form_data como lo arma extract_form_data: valor por id del mapping."""
    by_id = {}
    for fc in form.field_mapping:
        valor = valores_por_nombre.get(fc.get("name"))
        if valor is None:
            continue
        ids = fc.get("id") if isinstance(fc.get("id"), list) else [fc.get("id")]
        for i in ids:
            by_id[str(i)] = valor
    return {"__by_id": by_id}


def test_iphone_el_guion_omite_el_campo(iphone_brasil):
    datos = _form_data(iphone_brasil, {"Nombre": "-"})
    assert iphone_brasil._valor_para_etiqueta(datos, "texto", "Nome") is OMITIR


def test_iphone_el_excel_le_gana_al_id_dinamico(iphone_brasil):
    datos = _form_data(iphone_brasil, {"Apellido": "Pérez"})
    assert iphone_brasil._valor_para_etiqueta(datos, "texto", "Sobrenome") == "Pérez"


def test_iphone_celda_vacia_usa_el_id_dinamico(iphone_brasil):
    datos = _form_data(iphone_brasil, {"Apellido": ""})
    assert iphone_brasil._valor_para_etiqueta(datos, "texto", "Sobrenome") == "Fijo"


def test_iphone_select_vacio_va_al_azar(iphone_brasil):
    datos = _form_data(iphone_brasil, {"Modelo": ""})
    assert iphone_brasil._valor_para_etiqueta(datos, "select", "Modelos") is None


def test_iphone_select_con_dato_del_excel(iphone_brasil):
    datos = _form_data(iphone_brasil, {"Modelo": "S10"})
    assert iphone_brasil._valor_para_etiqueta(datos, "select", "Modelos") == "S10"


def test_iphone_cpf_recupera_el_cero_que_excel_se_come(iphone_brasil):
    cpf_col = next(fc["name"] for fc in iphone_brasil.field_mapping
                   if "cpf" in (fc.get("id") if isinstance(fc.get("id"), list) else [fc.get("id")]))
    datos = _form_data(iphone_brasil, {cpf_col: "9446875048"})
    assert iphone_brasil._valor_para_etiqueta(datos, "texto", "CPF") == "09446875048"


def test_iphone_campo_opcional_no_mapeado_se_decide_una_vez(iphone_brasil):
    datos = _form_data(iphone_brasil, {})
    primero = iphone_brasil._valor_para_etiqueta(datos, "texto", "Profesión")
    for _ in range(5):
        assert iphone_brasil._valor_para_etiqueta(datos, "texto", "Profesión") == primero


def test_iphone_campo_obligatorio_no_mapeado_siempre_se_completa(iphone_brasil):
    datos = _form_data(iphone_brasil, {})
    assert iphone_brasil._valor_para_etiqueta(datos, "texto", "Endereço*")


# --- BrowserStack respeta Consecutivo / Paralelo, igual que los navegadores ----------------

def test_la_nube_respeta_consecutivo_y_paralelo():
    """iPhone, Android y Mac corren uno tras otro en Consecutivo y a la vez en Paralelo,
    con la MISMA regla que Chrome/Firefox/Edge: _run_market no les da un trato aparte."""
    import ast
    import inspect
    import textwrap
    from osocio.interface import envio_leads
    fuente = textwrap.dedent(inspect.getsource(envio_leads))
    arbol = ast.parse(fuente)
    run_market = next(n for n in ast.walk(arbol)
                      if isinstance(n, ast.FunctionDef) and n.name == "_run_market")
    codigo = ast.get_source_segment(fuente, run_market)
    assert "_parallel = (excels_par or mercados_par) and len(sessions) > 1" in codigo
    assert "bs_dispositivo" not in codigo


# --- placeholders y checkboxes del iPhone (form de Chile, gm-envolve) -------------------

@pytest.mark.parametrize("texto", ["Seleccione", "Seleccionar", "Seleccioná", "SELECCIONE UNA OPCIÓN",
                                   "-- Elegí --", "Selecione", "Escolha", "Choose"])
def test_nunca_se_elige_un_placeholder(texto):
    from osocio.core.campos_inferencia import _is_placeholder
    assert _is_placeholder(texto)


def _nodo(tag, x, y, w, h, texto="", valor=None):
    import xml.etree.ElementTree as ET
    e = ET.Element("XCUIElementType" + tag, x=str(x), y=str(y), width=str(w), height=str(h), label=texto)
    if valor is not None:
        e.set("value", valor)
    return e


def test_checkbox_de_grupo_y_superindice_de_terminos():
    from osocio.providers.browserstack.ios_nativo import FormNativoIOS
    # Grupo: el Switch empieza en el título y termina en su opción.
    titulo = _nodo("StaticText", 30, 420, 200, 20, "PREFERENCIA DE CONTACTO")
    whatsapp = _nodo("StaticText", 74, 510, 90, 20, "WhatsApp")
    c = FormNativoIOS._checkbox(_nodo("Switch", 30, 420, 300, 112, valor="0"), [titulo, whatsapp])
    assert (c["grupo"], c["etiqueta"], c["fila"]) == ("PREFERENCIA DE CONTACTO", "WhatsApp", 520)
    # Términos: el "1" del superíndice (misma línea) no es el título de un grupo.
    texto = _nodo("StaticText", 119, 587, 260, 22, "Acepto términos y condiciones de registro")
    sup = _nodo("StaticText", 387, 587, 6, 10, "1")
    c = FormNativoIOS._checkbox(_nodo("Switch", 84, 587, 307, 25, valor="0"), [sup, texto])
    assert (c["grupo"], c["etiqueta"]) == ("", "Acepto términos y condiciones de registro")


def test_el_link_del_video_lleva_a_la_sesion_dentro_de_su_build(monkeypatch):
    import requests
    from osocio.providers.browserstack import bs_driver

    class _Resp:
        def json(self):
            return {"automation_session": {
                "video_url": "https://automate.browserstack.com/sessions/abc/video?token=x",
                "browser_url": "https://automate.browserstack.com/builds/b1/sessions/abc",
                "device": "", "os_version": "", "status": "done"}}
    monkeypatch.setattr(requests, "get", lambda *a, **k: _Resp())
    assert bs_driver.fetch_bs_session("abc", "u", "k")["video_url"] == \
        "https://automate.browserstack.com/builds/b1/sessions/abc"


# --- radios (form de OnStar Argentina, "¿Tu vehículo es Chevrolet?") ---------------------

def test_iphone_arma_una_pregunta_de_radio_con_sus_opciones():
    from osocio.providers.browserstack.ios_nativo import FormNativoIOS
    titulo = _nodo("StaticText", 36, 237, 195, 20, "¿Tu vehículo es Chevrolet?")
    patente = _nodo("StaticText", 46, 164, 51, 18, "Patente")
    si_txt, no_txt = _nodo("StaticText", 41, 305, 14, 18, "Si"), _nodo("StaticText", 94, 305, 20, 18, "No")
    si = _nodo("Other", 36, 267, 36, 58, valor="0")
    no = _nodo("Other", 79, 254, 49, 71, valor="1")
    estaticos = [patente, titulo, si_txt, no_txt]
    radios = [(e, FormNativoIOS._opcion_de_radio(e, estaticos)) for e in (si, no)]
    assert [_texto_de(s) for _, s in radios] == ["Si", "No"]
    [g] = FormNativoIOS._grupos_de_radios(radios, estaticos)
    assert (g["tipo"], g["etiqueta"], g["valor"]) == ("radio", "¿Tu vehículo es Chevrolet?", "No")
    assert [o["texto"] for o in g["opciones"]] == ["Si", "No"]
    # Un select (value = opción elegida) no es un radio.
    assert FormNativoIOS._opcion_de_radio(_nodo("Other", 36, 100, 300, 50, valor="Seleccionar"), estaticos) is None


def _texto_de(e):
    from osocio.providers.browserstack.ios_nativo import _texto
    return _texto(e)


class _Radio:
    def __init__(self, sel=False):
        self.sel = sel

    def is_selected(self):
        return self.sel


def _motor_radios(prefs=None):
    from osocio.core.form_filler.casillas import CasillasYRadiosMixin

    class _M(CasillasYRadiosMixin):
        def _pide_omitir_campo(self, fid, valor):
            from osocio.utils.omitir_campo import pide_omitir
            return pide_omitir(valor)
    m = _M()
    m.radio_prefs = prefs or {}
    m._decision_opcionales = {}
    return m


def _grupo(nombre, *valores, marcada=None):
    return {"name": nombre, "original_name": nombre,
            "options": [{"value": v, "index": i, "element": _Radio(v == marcada)} for i, v in enumerate(valores)]}


def test_radio_al_azar_varia_y_se_mantiene_en_la_fila():
    elegidas = set()
    for _ in range(40):
        m = _motor_radios()
        g = _grupo("have-chevrolet", "Si", "No")
        primera = m._choose_radio_option(g)["value"]
        assert m._choose_radio_option(g)["value"] == primera   # misma fila: misma opción
        elegidas.add(primera)
    assert elegidas == {"Si", "No"}
    # Aunque el form traiga una marcada de entrada, también se sortea.
    vistas = {_motor_radios()._choose_radio_option(_grupo("have-interest", "Renovar", "Suscribir",
                                                          marcada="Renovar"))["value"] for _ in range(40)}
    assert vistas == {"Renovar", "Suscribir"}


def test_radio_el_excel_manda_y_el_guion_lo_omite():
    assert _motor_radios({"have-chevrolet": "no"})._choose_radio_option(
        _grupo("have-chevrolet", "Si", "No"))["value"] == "No"
    assert _motor_radios({"have-chevrolet": "-"})._choose_radio_option(
        _grupo("have-chevrolet", "Si", "No")) is None
    # Libro de Reclamaciones: nunca "menor de edad" al azar.
    for _ in range(10):
        assert _motor_radios()._choose_radio_option(_grupo(
            "cc-younger-status", "Soy menor de edad", "Soy mayor de edad"))["value"] == "Soy mayor de edad"


# --- etiquetas de los forms con layout default (label afuera del campo) ---------------

@pytest.mark.parametrize("etiqueta, campo", [
    ("C.I.", "document"), ("CI *", "document"), ("CÉDULA", "document"), ("DNI", "document"),
    ("Ciudad", "city"), ("CIUDAD", "city"), ("MODELOS", "models"), ("FECHA ESTIMADA DE COMPRA", "estimated-day"),
])
def test_etiquetas_de_forms_con_layout_default(etiqueta, campo):
    assert campo in ids_para_etiqueta(etiqueta)


def test_browserstack_no_crea_carpeta_de_capturas_y_no_repite_numero(tmp_path):
    """En la nube la evidencia es el video: nunca carpeta screenshots_<...>N."""
    from osocio.core.generic_country_base import GenericCountryBase
    nube = GenericCountryBase("Colombia", browser="browserstack_mac", capturas=False)
    entrada = tmp_path / "entrada.xlsx"
    entrada.write_bytes(b"")
    nube.RESULTADOS_DIR, nube.EXCEL_PATH = str(tmp_path / "res"), str(entrada)
    os.makedirs(nube.RESULTADOS_DIR)
    nube.setup_directories_and_files()
    primero = nube.RESULTADOS_PATH
    nube.setup_directories_and_files()          # otra sesión, en paralelo, antes del Excel
    assert nube.RESULTADOS_PATH != primero
    assert not any(n.startswith("screenshots_") for n in os.listdir(nube.RESULTADOS_DIR))
    # Los navegadores locales siguen creando su carpeta.
    local = GenericCountryBase("Colombia", browser="chrome")
    local.RESULTADOS_DIR, local.EXCEL_PATH = nube.RESULTADOS_DIR, str(entrada)
    local.setup_directories_and_files()
    assert os.path.isdir(local.SCREENSHOT_DIR)


# --- Comparador de Dealers: URLs con espacio de más (aviso, no fallo) ---------------

def test_comparador_avisa_espacios_de_mas_sin_tratarlos_como_otra_url():
    from osocio.core.dealer_comparator_runner import _avisar_espacios_en_iframe, tiene_espacios_de_mas
    assert tiene_espacios_de_mas("https://x.com/form?model=onix ")
    assert tiene_espacios_de_mas(" https://x.com/form")
    assert not tiene_espacios_de_mas("https://x.com/form")
    assert not tiene_espacios_de_mas("")

    class _Iframe:
        def get_attribute(self, nombre):
            return {"data-src": "https://x.com/form?model=onix ", "src": ""}.get(nombre)
    avisos = []
    _avisar_espacios_en_iframe(_Iframe(), lambda msg, nivel="info": avisos.append((nivel, msg)))
    assert len(avisos) == 1 and avisos[0][0] == "warn" and "espacio de más" in avisos[0][1]
